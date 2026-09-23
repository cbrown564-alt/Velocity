from compute import *
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.table import Table, TableStyleInfo

F='Arial'; INK='1B2430'; MUTED='5B6573'; RULE='C9D1DB'; NAVY='1F3A5F'; TINT='EEF2F6'; NETF='F5F7FA'; SIGC='1F5FA8'; AMBER='B35C00'
def font(sz=10,b=False,c=INK,i=False,u=None): return Font(name=F,size=sz,bold=b,color=c,italic=i,underline=u)
fill=lambda c: PatternFill('solid',fgColor=c)
thin=Side(style='thin',color=RULE); med=Side(style='medium',color=NAVY)
PCT='0%;-0%;"–"'; NUM='#,##0'; MEAN='0.00'; NPSF='+0;-0;0'

df,meta=load(); tabs=build_tables(meta); groups=banner(df)
results=[run_table(t,df,groups) for t in tabs]
trend=trend_measures(df)
cur=df[df.wave==CUR]
wb=Workbook()

# ---------- Cover / Read me ----------
ws=wb.active; ws.title='Read me'; ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=110
lines=[('Atlas chilled-coffee brand tracker — Wave 4 data tables',font(16,True)),
 ('Reference tablebook generated from the synthetic brandtracker W1–W4 files for Velocity output-quality work (Track B).',font(10,c=MUTED)),
 ('',None),('About these tables',font(12,True,NAVY)),
 (f'Population: adults who bought ready-to-drink chilled coffee in the past 3 months ("category buyers").',font()),
 (f'Current wave: W4, fieldwork {FIELD[4]}. Unweighted sample 1,200. Previous wave for comparison: W3, fieldwork {FIELD[3]}, n=1,200.',font()),
 ('Weighting: rim weight wt (age × gender × region) in each wave. Weighted base sums to 1,200 per wave. Design effect ≈ 1.12 (effective sample ≈ 1,070).',font()),
 ('Percentages are of the stated base for each table. They are stored as numbers (0.73 displays as 73%) so they can be reused in calculations.',font()),
 ('',None),('How to read a table',font(12,True,NAVY)),
 ('Each column has a letter in brackets. A letter under a figure means that figure is significantly higher than the column with that letter, within the same banner group, at 95% confidence.',font()),
 ('Wave group: W4 (Total) is compared with W3. Segment, Age, Gender and Region columns are W4 respondents only, compared with each other within their group.',font()),
 ('Tests: two-sided z-test on weighted proportions (t-test for means) using effective bases (Kish). No correction for multiple comparisons.',font()),
 ('Rows are shown in scale or questionnaire order, not sorted by value. Nets (bold, shaded) combine the codes named in the label.',font()),
 ('',None),('Symbols and shorthand',font(12,True,NAVY)),
 ('–   true zero: no weighted respondents gave this answer.',font()),
 ('[u]   unweighted base below 30: estimate suppressed as unreliable. Not tested.',font()),
 ('Amber base   unweighted base 30–49: figures shown in italics and should be treated as indicative. Not tested.',font()),
 ('[x]   not applicable (for example, a mean for a non-scale question).',font()),
 ('',None),('Worksheets',font(12,True,NAVY)),
 ('Contents — every table with its base and a link.',font()),
 ('Key measures — headline brand measures for W4 with change since W3.',font()),
 ('Trend — the same measures for W1–W4.',font()),
 ('Chart data — the exact values plotted in the reference deck.',font()),
 ('T1 … — full tables with banner.',font()),
 ('Data (long) — every value in the tables in one tidy table: measure, group, wave, value, bases and status. Use it for pivots or checks.',font()),
 ('Definitions — questions, bases and derived variables.',font()),
 ('',None),('Source: synthetic data (brandtracker_w1–w4.sav). Brands and figures are fictitious.',font(9,c=MUTED,i=True))]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(i,1,t); 
    if f: c.font=f
    c.alignment=Alignment(wrap_text=True,vertical='top')

# ---------- Contents ----------
toc=wb.create_sheet('Contents'); toc.sheet_view.showGridLines=False
toc['A1']='Contents'; toc['A1'].font=font(16,True)
hdr=['Sheet','Section','Table title','Base','Unweighted base (W4)']
for j,h in enumerate(hdr,1):
    c=toc.cell(3,j,h); c.font=font(10,True,'FFFFFF'); c.fill=fill(NAVY); c.alignment=Alignment(vertical='center')
widths=[12,22,62,38,20]
for j,wd in enumerate(widths,1): toc.column_dimensions[CL(j)].width=wd
toc_rows=[('Key measures','Summary','Headline brand measures, W4 vs W3','All category buyers (conversion: aware of brand)',1200),
          ('Trend','Summary','Headline brand measures by wave, W1–W4','All category buyers (conversion: aware of brand)',1200),
          ('Chart data','Summary','Values plotted in the reference deck','As stated per chart',None)]
toc_rows+=[(f"T{r['t']['id']}",r['t']['section'],r['t']['title'],r['t']['basedesc'],r['bases'][1]['n']) for r in results]
toc_rows+=[('Data (long)','Data','All table values in long format','As stated per row',None),('Definitions','Data','Questions, bases and derived measures','',None)]
for i,row in enumerate(toc_rows,4):
    for j,v in enumerate(row,1):
        c=toc.cell(i,j,v); c.font=font(); c.border=Border(bottom=thin)
        if j==5 and v is not None: c.number_format=NUM
    c=toc.cell(i,1); c.hyperlink=f"#'{row[0]}'!A1"; c.font=font(10,c=SIGC,u='single')
toc.freeze_panes='A4'

# ---------- table sheets ----------
def write_table(r):
    t=r['t']; ws=wb.create_sheet(f"T{t['id']}"); ws.sheet_view.showGridLines=False
    cols=r['cols']; nc=len(cols)
    ws.column_dimensions['A'].width=40
    for j in range(nc): ws.column_dimensions[CL(j+2)].width=11.5
    ws['A1']='← Contents'; ws['A1'].hyperlink='#Contents!A1'; ws['A1'].font=font(9,c=SIGC,u='single')
    ws['A2']=f"Table {t['id']}. {t['title']}"; ws['A2'].font=font(13,True)
    ws['A3']=t['q']; ws['A3'].font=font(9,c=MUTED)
    ws['A4']=f"Base: {t['basedesc']}. W4 unless the column says otherwise. Weighted."; ws['A4'].font=font(9,c=MUTED)
    for rr in (3,4):
        ws.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=nc+1)
        ws.cell(rr,1).alignment=Alignment(wrap_text=True,vertical='top')
    ws.row_dimensions[3].height=13*max(1,-(-len(t['q'])//170))
    # banner rows 6-8
    R0=6
    g_start={}
    for j,c in enumerate(cols):
        col=j+2
        if c['group'] not in g_start:
            g_start[c['group']]=col
            gc=ws.cell(R0,col,c['group'])
        else: gc=ws.cell(R0,col)
        gc.fill=fill(NAVY); gc.font=font(10,True,'FFFFFF'); gc.alignment=Alignment(horizontal='centerContinuous')
        lc=ws.cell(R0+1,col,c['label']); lc.fill=fill(TINT); lc.font=font(9,True); lc.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        xc=ws.cell(R0+2,col,f"({c['letter']})"); xc.fill=fill(TINT); xc.font=font(9,c=MUTED); xc.alignment=Alignment(horizontal='center')
        left=Side(style='thin',color='FFFFFF') if g_start[c['group']]==col and col>2 else None
        if g_start[c['group']]==col and col>2:
            for rr in (R0,R0+1,R0+2): ws.cell(rr,col).border=Border(left=Side(style='medium',color='FFFFFF'))
    for rr in (R0,R0+1,R0+2): ws.cell(rr,1).fill=fill(NAVY if rr==R0 else TINT)
    ws.row_dimensions[R0+1].height=28
    # bases
    r_=R0+3
    for lab,key in (('Unweighted base','n'),('Weighted base','nw')):
        ws.cell(r_,1,lab).font=font(9,i=True,c=MUTED)
        for j,b in enumerate(r['bases']):
            c=ws.cell(r_,j+2,b[key]); c.number_format=NUM; c.alignment=Alignment(horizontal='right')
            if b['n']<MIN_SHOW: c.font=font(9,True,AMBER)
            elif b['n']<MIN_OK: c.font=font(9,True,AMBER); c.fill=fill('FFF1DE')
            else: c.font=font(9,i=True,c=MUTED)
        r_+=1
    for j in range(nc+1): ws.cell(r_-1,j+1).border=Border(bottom=Side(style='thin',color=INK))
    freeze_row=R0+3
    # body
    for row in r['rows']:
        isnet=row['kind'] in('net','score'); 
        lc=ws.cell(r_,1,row['label']); lc.font=font(10,isnet); lc.alignment=Alignment(indent=0 if isnet or row['kind']=='mean' else 1)
        for j,v in enumerate(row['vals']):
            n=r['bases'][j]['n']; c=ws.cell(r_,j+2)
            if n<MIN_SHOW:
                c.value='[u]'; c.font=font(9,c=MUTED)
            else:
                if row['kind']=='share': pass
                val=v['p']
                c.value=None if val!=val else float(val)
                c.number_format={'share':PCT,'mean':MEAN,'score':NPSF}[v['kind']]
                c.font=font(10,isnet,i=(n<MIN_OK))
            c.alignment=Alignment(horizontal='right')
            if isnet: c.fill=fill(NETF)
        if isnet: lc.fill=fill(NETF)
        r_+=1
        # sig row
        if any(v['sig'] for v in row['vals']):
            for j,v in enumerate(row['vals']):
                c=ws.cell(r_,j+2,v['sig'] or None); c.font=font(8,True,SIGC); c.alignment=Alignment(horizontal='right',vertical='top')
                if isnet: c.fill=fill(NETF)
            if isnet: ws.cell(r_,1).fill=fill(NETF)
            ws.row_dimensions[r_].height=11
            r_+=1
        for j in range(nc+1): ws.cell(r_-1,j+1).border=Border(bottom=thin)
    r_+=1
    notes=['Letters: significantly higher than the lettered column within the same banner group (95%, effective-base z-test). W4 (Total) is tested against W3.',
           '[u] base under 30, suppressed. Amber base 30–49: indicative only, not tested. – = true zero.']
    for nt in notes:
        ws.cell(r_,1,nt).font=font(8,c=MUTED); r_+=1
    ws.freeze_panes=ws.cell(freeze_row,2)
    ws.print_title_rows=f'{R0}:{R0+2}'; ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True

# ---------- Key measures ----------
km=wb.create_sheet('Key measures',2); km.sheet_view.showGridLines=False
km['A1']='Key measures — W4 (Q2 2026) compared with W3 (Q1 2026)'; km['A1'].font=font(14,True)
km['A2']=f'Base: all category buyers (W4 n=1,200; W3 n=1,200), except conversion (aware of brand). Weighted. Change in percentage points; ▲/▼ significant at 95%.'; km['A2'].font=font(9,c=MUTED)
heads=['Measure','Brand','W4','W3','Change (pts)','Significant?','W4 unweighted base']
for j,h in enumerate(heads,1):
    c=km.cell(4,j,h); c.font=font(10,True,'FFFFFF'); c.fill=fill(NAVY); c.alignment=Alignment(horizontal='left' if j<3 else 'right',wrap_text=True)
for j,wd in enumerate([44,12,10,10,13,13,20],1): km.column_dimensions[CL(j)].width=wd
rr=5; last=None
for m in trend:
    c4,c3=m['cells'][3],m['cells'][2]; ch=(c4['p']-c3['p'])*100; sig=m['p_prev']<ALPHA
    vals=[m['measure'] if m['measure']!=last else None,m['brand'],c4['p'],c3['p'],round(ch,1)+0.0,('▲ Yes' if ch>0 else '▼ Yes') if sig else 'No',c4['n']]
    for j,v in enumerate(vals,1):
        c=km.cell(rr,j,v); c.font=font(10,j==1 or (j==2 and m['brand']=='Atlas'))
        if j in(3,4): c.number_format='0.0%'
        if j==5: c.number_format='+0.0;-0.0;0.0'
        if j==6 and sig: c.font=font(10,True,'1E7B4F' if ch>0 else 'B3261E')
        if j==7: c.number_format=NUM
        if j>2: c.alignment=Alignment(horizontal='right')
        if m['brand']=='Atlas': c.fill=fill(TINT)
    if m['measure']!=last and last is not None:
        for j in range(1,8): km.cell(rr,j).border=Border(top=Side(style='thin',color=INK))
    last=m['measure']; rr+=1
km.freeze_panes='C5'

# ---------- Trend ----------
tr=wb.create_sheet('Trend',3); tr.sheet_view.showGridLines=False
tr['A1']='Key measures by wave, W1–W4'; tr['A1'].font=font(14,True)
tr['A2']='Base: all category buyers, except conversion (aware of brand). Weighted. ▲/▼ in a change column: that change is significant at 95%. Unweighted bases: 1,200 per wave for all-buyer measures.'; tr['A2'].font=font(9,c=MUTED)
heads=['Measure','Brand']+[WAVE_LBL[w] for w in (1,2,3,4)]+['Change W3→W4 (pts)','Change W1→W4 (pts)']
for j,h in enumerate(heads,1):
    c=tr.cell(4,j,h); c.font=font(10,True,'FFFFFF'); c.fill=fill(NAVY); c.alignment=Alignment(horizontal='left' if j<3 else 'right',wrap_text=True)
for j,wd in enumerate([44,12,12,12,12,12,16,16],1): tr.column_dimensions[CL(j)].width=wd
tr.row_dimensions[4].height=30
rr=5; last=None
for m in trend:
    ps=[c['p'] for c in m['cells']]
    d1=(ps[3]-ps[2])*100; d0=(ps[3]-ps[0])*100
    vals=[m['measure'] if m['measure']!=last else None,m['brand']]+ps+[round(d1,1)+0.0,round(d0,1)+0.0]
    for j,v in enumerate(vals,1):
        c=tr.cell(rr,j,v); c.font=font(10,j==1 or (j==2 and m['brand']=='Atlas'))
        if 3<=j<=6: c.number_format='0.0%'
        if j>=7: c.number_format='+0.0;-0.0;0.0'
        if j>2: c.alignment=Alignment(horizontal='right')
        if m['brand']=='Atlas': c.fill=fill(TINT)
    for j,key in ((7,'p_prev'),(8,'p_first')):
        if m[key]<ALPHA:
            c=tr.cell(rr,j); up=(d1 if j==7 else d0)>0
            c.number_format=('"▲ "+0.0;"▼ "-0.0') ; c.font=font(10,True,'1E7B4F' if up else 'B3261E')
    if m['measure']!=last and last is not None:
        for j in range(1,9): tr.cell(rr,j).border=Border(top=Side(style='thin',color=INK))
    last=m['measure']; rr+=1
tr.freeze_panes='C5'

wb._tmp=dict(results=results,trend=trend)
import pickle
for r in results: write_table(r)


# ---------- Chart data ----------
import json as _j
D=_j.load(open('deck.json'))
cd=wb.create_sheet('Chart data',4); cd.sheet_view.showGridLines=False
cd['A1']='Chart data — values plotted in Atlas_W4_readout_reference.pptx'; cd['A1'].font=font(14,True)
cd['A2']='One row per plotted value. Percentages as proportions; NPS in points. Change is W4 minus W3 in percentage points, from unrounded values.'; cd['A2'].font=font(9,c=MUTED)
H=['Slide','Exhibit','Series','Category','Value','Unit','Change vs W3 (pts)','Significant vs W3','Unweighted base','Base']
cd.append([]); cd.append(H)
for j,h in enumerate(H,1):
    c=cd.cell(4,j); c.font=font(10,True,'FFFFFF'); c.fill=fill(NAVY)
BN={b:BRAND_NAMES[b] for b in BR}; WL=['W1 Q3 2025','W2 Q4 2025','W3 Q1 2026','W4 Q2 2026']
rows=[]
for k,lab in [('unaided','Unaided awareness'),('aware','Aided awareness'),('consider','Would consider'),('p3m','Bought in past 3 months'),('pref','Preferred brand')]:
    for b in ('atlas','beacon','meridian'):
        m=D[k][b]; rows.append([3,'Brand funnel, W4',BN[b],lab,m['p'][3]/100,'proportion',m['ch'][3],'Yes' if m['sig'][3] else 'No',m['n'][3],'All category buyers'])
for k,lab in [('aware','Aided awareness'),('unaided','Unaided awareness')]:
    m=D[k]['atlas']
    for i in range(4): rows.append([4,'Atlas awareness by wave',lab,WL[i],m['p'][i]/100,'proportion',m['ch'][i],('Yes' if m['sig'][i] else 'No') if i else None,m['n'][i],'All category buyers'])
for b in ('beacon','meridian','atlas'):
    m=D['consider'][b]
    for i in range(4): rows.append([5,'Consideration by wave',BN[b],WL[i],m['p'][i]/100,'proportion',m['ch'][i],('Yes' if m['sig'][i] else 'No') if i else None,m['n'][i],'All category buyers'])
for b in ('beacon','solstice','atlas','cardinal','meridian'):
    m=D['conv'][b]
    for i in (2,3): rows.append([6,'Consideration among aware',WL[i],BN[b],m['p'][i]/100,'proportion',m['ch'][3] if i==3 else None,('Yes' if m['sig'][3] else 'No') if i==3 else None,m['n'][i],f'Aware of {BN[b]}'])
for b in ('atlas','beacon'):
    for a,v in D['image'][b].items(): rows.append([7,'Brand image, W4',BN[b],a,v['p']/100,'proportion',v['ch'],'Yes' if v['sig'] else 'No',v['n'],f'Aware of {BN[b]}'])
for r in D['seg']:
    for k,lab in (('aware','Aided'),('unaided','Unaided')):
        v=r[k]; rows.append([8,'Atlas awareness by group',lab,r['label'],v['w4']/100,'proportion',v['ch'],'Yes' if v['sig'] else 'No',v['n4'],'Category buyers in group'])
for n in D['nps']:
    for key,lab in (('det','Detractors (0–6)'),('pas','Passives (7–8)'),('pro','Promoters (9–10)')):
        rows.append([9,'Atlas NPS distribution',lab,WL[n['wave']-1],n[key]/100,'proportion',None,None,n['n'],'Bought/drank Atlas past 3 months'])
    rows.append([9,'Atlas NPS',"NPS",WL[n['wave']-1],n['nps'],'NPS points',None,('No' if n['wave']==4 else None),n['n'],'Bought/drank Atlas past 3 months'])
for r in rows:
    cd.append(r); rr=cd.max_row
    for j in range(1,11): cd.cell(rr,j).font=font(); cd.cell(rr,j).border=Border(bottom=thin)
    cd.cell(rr,5).number_format='0.0%' if r[5]=='proportion' else '+0.0;-0.0;0.0'
    cd.cell(rr,7).number_format='+0.0;-0.0;0.0'; cd.cell(rr,9).number_format=NUM
for j,wd in enumerate([7,30,22,30,10,12,14,14,12,30],1): cd.column_dimensions[CL(j)].width=wd
cd.freeze_panes='A5'

# ---------- Data (long) ----------
dl=wb.create_sheet('Data (long)'); 
cols=['table','table_title','question_base','row_label','row_type','banner_group','column','column_letter','wave','value','value_unit','unweighted_base','weighted_base','effective_base','status','sig_higher_than']
dl.append(cols)
for r in results:
    for row in r['rows']:
        for j,(c,v) in enumerate(zip(r['cols'],row['vals'])):
            b=r['bases'][j]; wave=PREV if c['label'].startswith('W3') else CUR
            status='suppressed_low_base' if b['n']<MIN_SHOW else ('indicative_low_base' if b['n']<MIN_OK else 'ok')
            val=None if status=='suppressed_low_base' or v['p']!=v['p'] else round(float(v['p']),6)
            unit={'share':'proportion','mean':'mean score','score':'NPS points'}[v['kind']]
            if status=='ok' and val==0 and v['kind']=='share': status='true_zero'
            dl.append([f"T{r['t']['id']}",r['t']['title'],r['t']['basedesc'],row['label'],row['kind'],c['group'],c['label'],c['letter'],f'W{wave}',val,unit,b['n'],round(b['nw'],2),round(b['neff'],1),status,v['sig'] or None])
for m in trend:
    for wv,c in zip((1,2,3,4),m['cells']):
        dl.append(['Trend',m['measure'],m['basedesc'],m['brand'],'measure','Wave',WAVE_LBL[wv],None,f'W{wv}',round(float(c['p']),6),'proportion',c['n'],round(c['nw'],2),round(c['neff'],1),'ok',None])
n=dl.max_row
tab=Table(displayName='TrackerLong',ref=f'A1:{CL(len(cols))}{n}'); tab.tableStyleInfo=TableStyleInfo(name='TableStyleLight1',showRowStripes=True)
dl.add_table(tab)
for j,wd in enumerate([8,40,30,34,9,12,16,8,6,10,14,10,10,10,20,10],1): dl.column_dimensions[CL(j)].width=wd
for row in dl.iter_rows(min_row=2):
    row[9].number_format='0.0000'
dl.freeze_panes='A2'

# ---------- Definitions ----------
de=wb.create_sheet('Definitions'); de.sheet_view.showGridLines=False
de['A1']='Definitions'; de['A1'].font=font(14,True)
rows=[('Measure','Definition','Base'),
 ('Unaided awareness','Brand mentioned at Q1/Q1a without prompting (any mention).','All category buyers'),
 ('Aided awareness','Brand selected at Q2 from a list.','All category buyers'),
 ('Consideration (all buyers)','Q3 "definitely" or "probably" consider. Not-aware respondents are not asked and count as not considering.','All category buyers'),
 ('Conversion (consideration among aware)','Q3 "definitely" or "probably" consider.','Aware of the brand (Q2)'),
 ('Past-3-month usage','Q4 bought or drunk in the past 3 months.','All category buyers'),
 ('Preference','Q5 single most preferred brand.','All category buyers'),
 ('Net Promoter Score','% scoring 9–10 minus % scoring 0–6 at Q6, in points.','Bought/drunk Atlas in past 3 months'),
 ('Image agreement','Q7 "strongly" or "somewhat" agree. Don\'t know counted in the base.','Aware of the brand'),
 ('Mean score','Weighted mean of scale codes, excluding don\'t know.','As table'),
 ('Effective base','(Σw)² / Σw²; used for all significance tests.','As table'),
 ('Segment','Consumer segment from the segmentation model (SEG).','All category buyers')]
for i,row in enumerate(rows,3):
    for j,v in enumerate(row,1):
        c=de.cell(i,j,v); c.font=font(10,i==3,'FFFFFF' if i==3 else INK); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
        if i==3: c.fill=fill(NAVY)
for j,wd in enumerate([34,80,34],1): de.column_dimensions[CL(j)].width=wd

import pickle; pickle.dump(dict(trend=trend),open('trend.pkl','wb'))
wb.save('Atlas_W4_tables_reference.xlsx'); print('saved', len(wb.sheetnames), wb.sheetnames[:8])
