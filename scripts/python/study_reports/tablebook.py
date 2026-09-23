"""Write a tablebook (.xlsx) for any canonical study from its tables, banner and summary sheets."""
from __future__ import annotations
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.table import Table as XTable, TableStyleInfo
from engine import share, mean, nps, column_tests, Est, _base

F = "Arial"; INK = "1B2430"; MUTED = "5B6573"; RULE = "C9D1DB"; NAVY = "1F3A5F"; TINT = "EEF2F6"; NETF = "F5F7FA"
SIGC = "1F5FA8"; AMBER = "B35C00"; AMBERF = "FFF1DE"; UP = "1E7B4F"; DOWN = "B3261E"
PCT = '0%;-0%;"–"'; NUM = "#,##0"


def font(sz=10, b=False, c=INK, i=False, u=None): return Font(name=F, size=sz, bold=b, color=c, italic=i, underline=u)
def fill(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=RULE)


def compute(st, table, banner, low):
    """Numbers for one table: per column base and per row estimates + letters."""
    w = st.w; cols = banner.columns
    bases = [Est(0, *_base(w, table.base & c.mask), float("nan")) for c in cols]
    out = []
    for r in table.rows:
        if r.kind == "excluded":
            out.append({"row": r, "counts": [int(((r.base & c.mask) & r.num).sum()) for c in cols]}); continue
        rb = r.base if r.base is not None else table.base
        if r.kind in ("share", "net"): ests = [share(r.num, rb & c.mask, w) for c in cols]
        elif r.kind == "mean": ests = [mean(r.x, rb & c.mask, w) for c in cols]
        else: ests = [nps(r.x, rb & c.mask, w) for c in cols]
        letters = column_tests(ests, banner, low["indicative"])
        out.append({"row": r, "ests": ests, "letters": letters})
    return bases, out


class Writer:
    def __init__(self, st, spec, banner):
        self.st, self.spec, self.banner = st, spec, banner
        self.low = spec.get("low_base", {"suppress": 30, "indicative": 50})
        self.wb = Workbook(); self.long = []; self.contents = []; self.summaries = []

    # ------------------------------------------------------------------ front matter
    def readme(self, extra_lines=()):
        ws = self.wb.active; ws.title = "Read me"; ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 120
        cb, sp = self.st.cb, self.spec
        L = [(f"{sp['report_title']} — data tables", font(16, True)), (sp.get("source_note", ""), font(10, c=MUTED)), ("", None),
             ("About this study", font(12, True, NAVY))]
        for k, lab in (("population", "Population"), ("decision_question", "Decision question")):
            if cb.get(k): L.append((f"{lab}: {cb[k]}", font()))
        d = cb.get("design", {})
        L.append((f"Design: {d.get('type', '').replace('_', ' ')}. Unweighted sample {len(self.st.df):,}.", font()))
        if cb.get("waves"): L.append(("Waves: " + ", ".join(w["label"] for w in cb["waves"]) + ". Each wave is a fresh sample.", font()))
        L.append((f"Weight: {cb['weight']} (all percentages are weighted).", font()))
        L += [(x, font()) for x in extra_lines]
        L += [("", None), ("How to read a table", font(12, True, NAVY)),
              ("Base: respondents who were asked the question and gave a valid answer. Don't know, prefer not to say and not answered are shown as counts below each table and are not in the base.", font()),
              ("Each column has a letter. A letter under a figure means that figure is significantly higher than the lettered column in the same group (95%, two-sided Wald test on Kish effective bases).", font())]
        for g in self.banner.groups():
            c0 = next(c for c in self.banner.columns if c.group == g)
            if c0.compare_prev_only: L.append((f"{g} group: each column is compared only with the column before it (wave on wave).", font()))
            if self.banner.adjust.get(g) == "holm": L.append((f"{g} group: pairwise tests are Holm-adjusted within each row.", font()))
        L += [(f"Unweighted base under {self.low['suppress']}: figures suppressed as [u]. Base {self.low['suppress']}–{self.low['indicative'] - 1}: figures in italics with an amber base, indicative only and not tested.", font()),
              ("– = true zero. Rows are in questionnaire order. Nets are bold and shaded.", font()), ("", None),
              ("Missing-value convention", font(12, True, NAVY))]
        for k, v in cb["missing_policy"].items(): L.append((f"{'blank' if k == 'sysmis' else k}: {v}", font()))
        L += [("", None), ("Sheets", font(12, True, NAVY)), ("Contents — every table with its base and a link. Key measures — headline measures. Chart data — values plotted in the deck. "
              "T sheets — full tables. Data (long) — every value in one tidy table. Definitions — every variable, its question, universe and codes.", font())]
        for i, (t, f) in enumerate(L, 1):
            c = ws.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
            if f: c.font = f

    # ------------------------------------------------------------------ key measures (archetype-provided grid)
    def grid_sheet(self, name, grid, index=None):
        ws = self.wb.create_sheet(name, index) if index is not None else self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["A1"] = grid["title"]; ws["A1"].font = font(14, True)
        ws["A2"] = grid["subtitle"]; ws["A2"].font = font(9, c=MUTED)
        heads = grid["columns"]
        for j, h in enumerate(heads, 1):
            c = ws.cell(4, j, h); c.font = font(10, True, "FFFFFF"); c.fill = fill(NAVY)
            c.alignment = Alignment(horizontal="left" if j <= grid.get("label_cols", 2) else "right", wrap_text=True, vertical="bottom")
        ws.row_dimensions[4].height = 30
        for j, wd in enumerate(grid.get("widths", []), 1): ws.column_dimensions[CL(j)].width = wd
        r = 5; last = None
        for row in grid["rows"]:
            first_in_group = row.get("group") != last
            for j, v in enumerate(row["cells"], 1):
                c = ws.cell(r, j, v["v"] if isinstance(v, dict) else v)
                st_ = v if isinstance(v, dict) else {}
                c.font = font(10, st_.get("bold", j == 1 and first_in_group), st_.get("color", INK), st_.get("italic", False))
                if st_.get("fmt"): c.number_format = st_["fmt"]
                if j > grid.get("label_cols", 2): c.alignment = Alignment(horizontal="right")
                if row.get("highlight"): c.fill = fill(TINT)
                if first_in_group and last is not None: c.border = Border(top=Side(style="thin", color=INK))
            last = row.get("group"); r += 1
        ws.freeze_panes = ws.cell(5, grid.get("label_cols", 2) + 1)
        self.summaries.append((name, "Summary", grid["title"], grid.get("base", ""), None))

    # ------------------------------------------------------------------ tables
    def table_sheet(self, table):
        st, banner, low = self.st, self.banner, self.low
        bases, rows = compute(st, table, banner, low)
        ws = self.wb.create_sheet(table.id); ws.sheet_view.showGridLines = False
        cols = banner.columns; nc = len(cols)
        ws.column_dimensions["A"].width = 42
        for j in range(nc): ws.column_dimensions[CL(j + 2)].width = 10.5
        ws["A1"] = "← Contents"; ws["A1"].hyperlink = "#Contents!A1"; ws["A1"].font = font(9, c=SIGC, u="single")
        ws["A2"] = f"Table {table.id[1:]}. {table.title}"; ws["A2"].font = font(13, True)
        ws["A3"] = table.question; ws["A3"].font = font(9, c=MUTED)
        scope = self.spec.get("banner_scope_text")
        ws["A4"] = f"Base: {table.base_text}. Weighted." + (f" {scope}" if scope else ""); ws["A4"].font = font(9, c=MUTED)
        for rr in (3, 4):
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=nc + 1)
            ws.cell(rr, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[3].height = 13 * max(1, -(-len(table.question) // 190))
        R0 = 6; seen = set()
        for j, c in enumerate(cols):
            col = j + 2; first = c.group not in seen; seen.add(c.group)
            g = ws.cell(R0, col, c.group if first else None); g.fill = fill(NAVY); g.font = font(10, True, "FFFFFF"); g.alignment = Alignment(horizontal="centerContinuous")
            l = ws.cell(R0 + 1, col, c.label); l.fill = fill(TINT); l.font = font(9, True); l.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            x = ws.cell(R0 + 2, col, f"({c.letter})"); x.fill = fill(TINT); x.font = font(9, c=MUTED); x.alignment = Alignment(horizontal="center")
            if first and col > 2:
                for rr in (R0, R0 + 1, R0 + 2): ws.cell(rr, col).border = Border(left=Side(style="medium", color="FFFFFF"))
        for rr in (R0, R0 + 1, R0 + 2): ws.cell(rr, 1).fill = fill(NAVY if rr == R0 else TINT)
        ws.row_dimensions[R0 + 1].height = 30
        r = R0 + 3
        for lab, key in (("Unweighted base", "n"), ("Weighted base", "nw")):
            ws.cell(r, 1, lab + (" (table)" if table.row_bases_vary else "")).font = font(9, i=True, c=MUTED)
            for j, b in enumerate(bases):
                v = b.n if key == "n" else b.nw
                c = ws.cell(r, j + 2, "varies" if table.row_bases_vary else v); c.number_format = NUM; c.alignment = Alignment(horizontal="right")
                if not table.row_bases_vary and b.n < low["indicative"]:
                    c.font = font(9, True, AMBER); c.fill = fill(AMBERF) if b.n >= low["suppress"] else fill("FFFFFF")
                else: c.font = font(9, i=True, c=MUTED)
            r += 1
        for j in range(nc + 1): ws.cell(r - 1, j + 1).border = Border(bottom=Side(style="thin", color=INK))
        freeze = r
        excluded = []
        for item in rows:
            row = item["row"]
            if row.kind == "excluded": excluded.append(item); continue
            isnet = row.kind in ("net", "nps")
            lc = ws.cell(r, 1, row.label); lc.font = font(10, isnet); lc.alignment = Alignment(indent=0 if isnet or row.kind == "mean" else 1, wrap_text=True)
            for j, e in enumerate(item["ests"]):
                c = ws.cell(r, j + 2); n = e.n
                if n < low["suppress"]:
                    c.value = "[u]"; c.font = font(9, c=MUTED)
                else:
                    c.value = None if e.value != e.value else float(e.value)
                    c.number_format = {"share": PCT, "mean": "0.00", "nps": "+0;-0;0"}[e.kind]
                    c.font = font(10, isnet, i=n < low["indicative"])
                c.alignment = Alignment(horizontal="right")
                if isnet: c.fill = fill(NETF)
                col = cols[j]
                self.long.append([table.id, table.title, table.base_text, row.label, row.kind, col.group, col.label, col.letter,
                                  None if n < low["suppress"] or e.value != e.value else round(float(e.value), 6),
                                  {"share": "proportion", "mean": "mean", "nps": "NPS points"}[e.kind], n, round(e.nw, 2), round(e.ess, 1),
                                  "suppressed_low_base" if n < low["suppress"] else "indicative_low_base" if n < low["indicative"] else ("true_zero" if e.value == 0 else "ok"),
                                  item["letters"][j] or None])
            if isnet: lc.fill = fill(NETF)
            r += 1
            if any(item["letters"]):
                for j, L_ in enumerate(item["letters"]):
                    c = ws.cell(r, j + 2, L_ or None); c.font = font(8, True, SIGC); c.alignment = Alignment(horizontal="right", vertical="top")
                    if isnet: c.fill = fill(NETF)
                if isnet: ws.cell(r, 1).fill = fill(NETF)
                ws.row_dimensions[r].height = 11; r += 1
            for j in range(nc + 1): ws.cell(r - 1, j + 1).border = Border(bottom=thin)
        if table.row_bases_vary:
            for item in rows:
                if item["row"].kind == "excluded": continue
                ws.cell(r, 1, f"Base: {item['row'].label}").font = font(8, i=True, c=MUTED)
                for j, e in enumerate(item["ests"]):
                    c = ws.cell(r, j + 2, e.n); c.number_format = NUM; c.alignment = Alignment(horizontal="right")
                    c.font = font(8, e.n < low["indicative"], AMBER if e.n < low["indicative"] else MUTED, i=True)
                r += 1
        for item in excluded:
            ws.cell(r, 1, item["row"].label + " — unweighted count").font = font(8, i=True, c=MUTED)
            for j, k in enumerate(item["counts"]):
                c = ws.cell(r, j + 2, k); c.number_format = NUM; c.font = font(8, i=True, c=MUTED); c.alignment = Alignment(horizontal="right")
            r += 1
        r += 1
        for nt in ["Letters: significantly higher than the lettered column in the same group (95%). See Read me for comparison rules.",
                   f"[u] unweighted base under {low['suppress']}, suppressed. Amber base {low['suppress']}–{low['indicative'] - 1}: indicative, not tested. – = true zero."] + table.notes:
            ws.cell(r, 1, nt).font = font(8, c=MUTED); r += 1
        ws.freeze_panes = ws.cell(freeze, 2)
        ws.print_title_rows = f"{R0}:{R0 + 2}"; ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        total_n = bases[0].n if not table.row_bases_vary else None
        self.contents.append((table.id, table.section, table.title, table.base_text, total_n))

    # ------------------------------------------------------------------ back matter
    def finish(self, path, chart_rows=None):
        wb = self.wb
        if chart_rows:
            cd = wb.create_sheet("Chart data", 1 + len(self.summaries))
            cd.sheet_view.showGridLines = False
            cd["A1"] = "Chart data — values plotted in the deck"; cd["A1"].font = font(14, True)
            cd["A2"] = "One row per plotted value. Proportions as 0–1; NPS in points; means as scored."; cd["A2"].font = font(9, c=MUTED)
            H = ["Slide", "Exhibit", "Series", "Category", "Value", "Unit", "Unweighted base", "Note"]
            for j, h in enumerate(H, 1):
                c = cd.cell(4, j, h); c.font = font(10, True, "FFFFFF"); c.fill = fill(NAVY)
            for i, row in enumerate(chart_rows, 5):
                for j, v in enumerate(row, 1):
                    c = cd.cell(i, j, v); c.font = font(); c.border = Border(bottom=thin)
                cd.cell(i, 5).number_format = "0.0%" if row[5] == "proportion" else "0.0"
                cd.cell(i, 7).number_format = NUM
            for j, wd in enumerate([7, 40, 24, 30, 10, 12, 14, 40], 1): cd.column_dimensions[CL(j)].width = wd
            cd.freeze_panes = "A5"
            self.summaries.append(("Chart data", "Summary", "Values plotted in the deck", "As stated per chart", None))
        # contents
        toc = wb.create_sheet("Contents", 1); toc.sheet_view.showGridLines = False
        toc["A1"] = "Contents"; toc["A1"].font = font(16, True)
        for j, h in enumerate(["Sheet", "Section", "Title", "Base", "Unweighted base (first column)"], 1):
            c = toc.cell(3, j, h); c.font = font(10, True, "FFFFFF"); c.fill = fill(NAVY)
        for j, wd in enumerate([13, 22, 58, 52, 16], 1): toc.column_dimensions[CL(j)].width = wd
        for i, row in enumerate(self.summaries + self.contents, 4):
            for j, v in enumerate(row, 1):
                c = toc.cell(i, j, v); c.font = font(); c.border = Border(bottom=thin)
                if j == 5 and v is not None: c.number_format = NUM
            c = toc.cell(i, 1); c.hyperlink = f"#'{row[0]}'!A1"; c.font = font(10, c=SIGC, u="single")
        toc.freeze_panes = "A4"
        # long
        dl = wb.create_sheet("Data (long)")
        H = ["table", "table_title", "base", "row_label", "row_type", "banner_group", "column", "column_letter", "value", "value_unit",
             "unweighted_base", "weighted_base", "effective_base", "status", "sig_higher_than"]
        dl.append(H)
        for row in self.long: dl.append(row)
        t = XTable(displayName="LongData", ref=f"A1:{CL(len(H))}{dl.max_row}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        dl.add_table(t); dl.freeze_panes = "A2"
        for j, wd in enumerate([7, 36, 30, 34, 8, 14, 16, 7, 10, 12, 10, 10, 10, 18, 10], 1): dl.column_dimensions[CL(j)].width = wd
        # definitions
        de = wb.create_sheet("Definitions"); de.sheet_view.showGridLines = False
        de["A1"] = "Definitions"; de["A1"].font = font(14, True)
        H = ["Variable", "Question ID", "Label", "Question text", "Wording source", "Universe (who was asked)", "Codes", "Missing codes"]
        for j, h in enumerate(H, 1):
            c = de.cell(3, j, h); c.font = font(10, True, "FFFFFF"); c.fill = fill(NAVY)
        for i, v in enumerate(self.st.cb["variables"], 4):
            vals = "; ".join(f"{x['code']} = {x['label']}" for x in v["values"])
            row = [v["name"], v.get("question_id"), v["label"], v.get("question_text"), v.get("text_source"), v["universe"]["text"],
                   vals[:500], ", ".join(str(m) for m in v["missing"])]
            for j, x in enumerate(row, 1):
                c = de.cell(i, j, x); c.font = font(9, c=AMBER if (j == 5 and x == "authored") else INK); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
        for j, wd in enumerate([24, 10, 34, 50, 12, 36, 50, 10], 1): de.column_dimensions[CL(j)].width = wd
        de.freeze_panes = "B4"
        wb.save(path)
