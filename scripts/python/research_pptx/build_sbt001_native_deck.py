"""Generate the SBT-001 Category Board using native Office charts for all data marks."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
sys.path.insert(0,str(Path(__file__).resolve().parent))
from native_chart_system import NativeChartBuilder,ChartSpec,Primitive,ResearchMeta

OUT=Path(__file__).resolve().parents[3]/'evals/research_quality/projects/synthetic/SBT-001/reference/deck';OUT.mkdir(parents=True,exist_ok=True)
BG=RGBColor(248,246,241);INK=RGBColor(30,38,43);MUTED=RGBColor(105,112,113);ACCENT=RGBColor(0,126,125);NAVY=RGBColor(35,55,72);RED=RGBColor(184,65,55);AMBER=RGBColor(184,126,43);WHITE=RGBColor(255,255,255)
W=Inches(13.333);H=Inches(7.5)
BR={'Northstar':(46.68,45.07),'Pulse':(20.29,29.59),'Mosaic':(22.13,25.69),'Lumen':(18.05,18.96),'Harbour':(23.66,24.12)}
PULSE=[19.49,20.29,27.21,29.07,29.59];HT=[25.07,23.66,23.71,22.79,24.12];HY=[22.06,19.47,17.73,14.95,15.02]

def text(s,t,x,y,w,h,size=18,color=INK,bold=False):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;p=tf.paragraphs[0];r=p.add_run();r.text=t;r.font.name='Aptos';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color;return q

def background(s,color=BG):
 q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,W,H);q.fill.solid();q.fill.fore_color.rgb=color;q.line.fill.background();s.shapes._spTree.remove(q._element);s.shapes._spTree.insert(2,q._element)

def headline(s,h,dek=None):
 text(s,h,.7,.42,12,.82,26,INK,True)
 if dek:text(s,dek,.72,1.25,11.7,.38,15,MUTED)

def build(path=OUT/'SBT001_Category_Board_native.pptx'):
 prs=Presentation();prs.slide_width=W;prs.slide_height=H;blank=prs.slide_layouts[6]
 # 1
 s=prs.slides.add_slide(blank);background(s);text(s,'CATEGORY BOARD',.72,.62,3,.25,11,ACCENT,True);text(s,'Pulse advances;\nNorthstar and Harbour face\ndifferent vulnerabilities.',.72,1.22,10.8,2.4,36,INK,True);text(s,'SBT-001 synthetic UK mobile brand tracker · Waves 1–5',.75,5.92,8.5,.35,17,MUTED)
 # 2
 s=prs.slides.add_slide(blank);background(s);headline(s,'Three pressures are reshaping the category.','Three distinct measures, populations and periods — each remains an editable native chart.')
 specs=[ChartSpec(Primitive.SLOPE_TWO_PERIOD,[{'category':'Pulse','start':20.29,'end':27.21}],y_min=15,y_max=30),ChartSpec(Primitive.TREND_HERO,[{'period':'W3','value':-12.54},{'period':'W4','value':-60.71},{'period':'W5','value':-33.17}],y_min=-70,y_max=0,series_labels=['Northstar NPS']),ChartSpec(Primitive.SLOPE_TWO_PERIOD,[{'category':'Harbour 18–34','start':22.06,'end':15.02}],y_min=10,y_max=25)]
 for i,sp in enumerate(specs):NativeChartBuilder(s).add(sp,.6+i*4.15,2.0,3.8,3.8)
 text(s,'PULSE · total-sample consideration',.7,5.95,3.5,.25,10,ACCENT,True);text(s,'NORTHSTAR · current-customer NPS',4.85,5.95,3.5,.25,10,RED,True);text(s,'HARBOUR · 18–34 consideration',9,5.95,3.5,.25,10,MUTED,True)
 # 3
 s=prs.slides.add_slide(blank);background(s);headline(s,'Northstar leads on consideration; Pulse records the largest gain.','W5 total-sample consideration × W2→W5 percentage-point change');rows=[{'entity':k,'position':v[1],'change':v[1]-v[0]} for k,v in BR.items()];NativeChartBuilder(s).add(ChartSpec(Primitive.SCATTER_POSITION_CHANGE,rows,x_min=10,x_max=50,y_min=-4,y_max=12),.9,1.75,11.5,4.9)
 # 4
 s=prs.slides.add_slide(blank);background(s);headline(s,"Pulse's consideration gains persist beyond the Wave 3 jump.",'Weighted total-sample consideration · campaign timing is contextual, not causal');NativeChartBuilder(s).add(ChartSpec(Primitive.TREND_HERO,[{'period':f'W{i+1}','value':v} for i,v in enumerate(PULSE)],y_min=15,y_max=32,series_labels=['Pulse consideration']),.9,1.75,11.5,4.9)
 # 5
 s=prs.slides.add_slide(blank);background(s);headline(s,"Pulse's observed consideration increase is larger among 18–34s.",'Native XY dumbbell · W2→W3 weighted consideration');NativeChartBuilder(s).add(ChartSpec(Primitive.DUMBBELL_XY,[{'category':'18–34','start':25.98,'end':37.12},{'category':'35+','start':18.07,'end':23.33}],x_min=15,x_max=40),.9,1.75,11.5,4.9)
 # 6
 s=prs.slides.add_slide(blank);background(s);headline(s,"Pulse's consideration and preference gains outpace gains in provider share.",'Native multi-row XY dumbbell · four separate population measures');rows=[{'measure':'Prompted awareness','start':57.79,'end':73.82},{'measure':'Consideration','start':20.29,'end':29.59},{'measure':'Preference','start':15.34,'end':25.11},{'measure':'Current provider','start':15.16,'end':19.79}];NativeChartBuilder(s).add(ChartSpec(Primitive.DUMBBELL_XY_MULTI,rows,x_min=0,x_max=80),.9,1.75,11.5,4.9)
 # 7
 s=prs.slides.add_slide(blank);background(s);headline(s,"Northstar's awareness remains high while customer advocacy falls sharply.",'Separate native charts preserve incompatible units and universes.');NativeChartBuilder(s).add(ChartSpec(Primitive.TREND_HERO,[{'period':'W3','value':-12.54},{'period':'W4','value':-60.71}],y_min=-70,y_max=0,series_labels=['NPS']),.75,2,5.8,4.2);NativeChartBuilder(s).add(ChartSpec(Primitive.TREND_HERO,[{'period':'W3','value':93.54},{'period':'W4','value':93.57}],y_min=85,y_max=100,series_labels=['Awareness']),6.8,2,5.8,4.2);text(s,'CURRENT NORTHSTAR CUSTOMERS · NPS',.85,1.7,3.5,.25,10,RED,True);text(s,'ALL ELIGIBLE RESPONDENTS · AWARENESS',6.9,1.7,4,.25,10,NAVY,True)
 # 8
 s=prs.slides.add_slide(blank);background(s);headline(s,'Northstar improves in Wave 5, but remains well short of its pre-shock position.','Two independent native recovery charts.');NativeChartBuilder(s).add(ChartSpec(Primitive.TREND_HERO,[{'period':'W3','value':-12.54},{'period':'W4','value':-60.71},{'period':'W5','value':-33.17}],y_min=-70,y_max=0,series_labels=['NPS']),.75,2,5.8,4.2);NativeChartBuilder(s).add(ChartSpec(Primitive.TREND_HERO,[{'period':'W3','value':13.78},{'period':'W4','value':34.09},{'period':'W5','value':21.76}],y_min=0,y_max=40,series_labels=['Recent problems']),6.8,2,5.8,4.2);text(s,'NPS',.85,1.7,1,.25,10,RED,True);text(s,'RECENT PROBLEMS',6.9,1.7,2,.25,10,AMBER,True)
 # 9
 s=prs.slides.add_slide(blank);background(s);headline(s,"Harbour's aggregate result masks declining consideration among 18–34s.",'Weighted consideration · common scale · five waves');rows=[{'period':f'W{i+1}','series_1':HT[i],'series_2':HY[i]} for i in range(5)];NativeChartBuilder(s).add(ChartSpec(Primitive.TREND_COMPARE,rows,y_min=10,y_max=30,series_labels=['Aggregate','18–34']),.9,1.75,11.5,4.9)
 # 10
 s=prs.slides.add_slide(blank);background(s,NAVY);text(s,'THREE QUESTIONS SHOULD SHAPE THE NEXT WAVE',.75,.65,10.8,.35,12,RGBColor(210,221,224),True);text(s,'What should the next wave tell us?',.75,1.25,10.8,.75,34,WHITE,True)
 for i,(b,q) in enumerate([('PULSE','Are consideration gains translating into provider choice?'),('NORTHSTAR','Is the customer recovery continuing?'),('HARBOUR','Is younger-audience consideration declining or beginning to recover?')]):text(s,b,.85,2.55+i*1.15,1.5,.25,11,ACCENT,True);text(s,q,2.55,2.5+i*1.15,8.5,.45,18,WHITE,True)
 prs.save(path);return path

if __name__=='__main__':print(build())
